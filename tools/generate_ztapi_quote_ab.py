"""Convert Yunxin's three-sheet quote into an A/B-only audit manifest."""

import hashlib
import json
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl


TOKEN_PRICE_KEYS = {
    "输入单价": "input_tokens",
    "输出单价": "output_tokens",
    "输出": "output_tokens",
    "缓存命中": "cache_read",
    "缓存写入": "cache_write",
    "缓存写入5m": "cache_write_5m",
    "缓存写入1h": "cache_write_1h",
}

# The 9.15 workbook's gateway row for GPT 5.4 Nano (F89) repeats GPT 5.6 Terra's
# price; its original-vendor row (F63) carries the real official price.
OFFICIAL_PRICE_RESOURCE_OVERRIDES = {
    ("gpt 5.4 nano", "网关"): "原厂",
}

MEDIA_MODELS = {
    "GPT Image 2": "image",
    "Gemini 2.5 Flash Image": "image",
    "Seedance 2.0": "video",
    "Seedance 2.0 Fast": "video",
    "Seedance 2.0 Mini": "video",
    "Seedance 2.5": "video",
}


def _format_price(value):
    return format(value.quantize(Decimal("0.0000000001"), rounding=ROUND_HALF_UP).normalize(), "f")


def _read_money(raw):
    value = raw.strip().replace("US$", "$")
    if value.startswith("$"):
        return "USD", Decimal(value[1:])
    if value.startswith("¥") or value.startswith("￥"):
        return "CNY", Decimal(value[1:])
    raise ValueError(f"unsupported quotation price {raw!r}")


def parse_token_price_rules(text, quoted_fraction, cost_share):
    """Parse only token price dimensions; unrecognized monetary fields fail closed."""
    if not quoted_fraction > 0 or not cost_share > 0:
        raise ValueError("quotation share must be positive")
    rules = []
    segments = text.rstrip("；").split("；")
    for segment in segments:
        segment = segment.strip()
        if not segment:
            raise ValueError("empty quotation tier")
        if re.fullmatch(r"输入\s*(?:US)?\$[0-9.]+", segment):
            segment = "输入单价=" + re.sub(r"^输入\s*", "", segment)
        values = [part.split("=", 1)[1].strip() for part in segment.split("｜") if "=" in part]
        if values and all(value == "-" for value in values):
            # A tier priced "-" everywhere is not offered; requests beyond it are rejected.
            continue
        currency = None
        prices = {}
        conditions = []
        not_applicable = []
        temporary_free = []
        for part in segment.split("｜"):
            part = part.strip()
            if "=" not in part:
                if part.startswith(("输入长度", "输出长度")):
                    conditions.append(part)
                    continue
                raise ValueError(f"unrecognized quotation fragment {part!r}")
            key, value = (item.strip() for item in part.split("=", 1))
            if key == "缓存存储":
                if value != "限时免费":
                    raise ValueError(f"paid cache storage needs separate metering: {value!r}")
                temporary_free.append("cache_storage")
                continue
            dimension = TOKEN_PRICE_KEYS.get(key)
            if dimension is None:
                if key in ("输入长度", "输出长度", "计费时段", "输入类型"):
                    conditions.append(part)
                    continue
                raise ValueError(f"unrecognized quotation field {key!r}")
            if value == "不适用":
                not_applicable.append(dimension)
                continue
            unit, amount = _read_money(value)
            if currency is not None and unit != currency:
                raise ValueError("mixed currencies in a single quotation tier")
            if dimension in prices or amount < 0:
                raise ValueError(f"duplicate or negative quotation price for {dimension}")
            currency = unit
            prices[dimension] = amount
        if currency is None or "input_tokens" not in prices:
            raise ValueError("token quotation tier lacks input price")
        rules.append({
            "conditions": conditions,
            "currency": currency,
            "not_applicable": not_applicable,
            "temporary_free": temporary_free,
            "cost": {key: _format_price(value * quoted_fraction) for key, value in prices.items()},
            "sale": {key: _format_price(value * quoted_fraction / cost_share) for key, value in prices.items()},
        })
    return rules


def normalize_name(value):
    name = str(value).strip()
    name = re.sub(r"\s*[\uFF08(](?:hc|os)[\uFF09)]\s*$", "", name, flags=re.I)
    return re.sub(r"^Dreamina\s+", "", name, flags=re.I)


def main():
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    mapped = {}
    for index, row in enumerate(workbook["\u6a21\u578b\u6620\u5c04\u5173\u7cfb"].iter_rows(min_row=2), 2):
        grade = str(row[2].value)[0]
        key = (str(row[5].value).strip().casefold(), grade)
        mapped.setdefault(key, []).append((index, str(row[3].value), str(row[4].value)))

    official = {}
    for index, row in enumerate(workbook["\u5404\u6a21\u578b\u5b98\u65b9\u62a5\u4ef7"].iter_rows(min_row=2), 2):
        key = (str(row[3].value).strip().casefold(), str(row[2].value))
        value = (str(row[5].value), f"F{index}")
        if key in official and official[key] != value:
            raise ValueError(f"ambiguous official price for {key}")
        official[key] = value

    entries = []
    for index, row in enumerate(workbook["\u7f51\u6613\u62a5\u4ef7"].iter_rows(min_row=2), 2):
        grade = str(row[1].value)[0]
        if grade not in ("A", "B"):
            continue
        code = str(row[2].value).strip()
        candidates = mapped[(code.casefold(), grade)]
        if len(candidates) == 1:
            _, resource, original = candidates[0]
        else:
            matching = [candidate for candidate in candidates if candidate[0] == index]
            if len(matching) != 1:
                raise ValueError(f"ambiguous mapping for {(code, grade)}")
            _, resource, original = matching[0]
        price_resource = OFFICIAL_PRICE_RESOURCE_OVERRIDES.get((original.casefold(), resource), resource)
        price, price_cell = official[(original.casefold(), price_resource)]
        status = str(row[4].value or "")
        entries.append(
            {
                "model_name": normalize_name(original),
                "model_original": original,
                "model_code": code,
                "grade": grade,
                "resource_type": resource,
                "quoted_fraction": str(Decimal(str(row[3].value))),
                "official_price_text": price,
                "quotation_cell": f"D{index}",
                "official_price_cell": price_cell,
                "status": status,
                "active": "\u4e0b\u7ebf" not in status and "\u4e0b\u67b6" not in status,
            }
        )

    if len(entries) != 72 or len({entry["model_name"].casefold() for entry in entries}) != 50:
        raise ValueError("A/B quotation row or model count changed")
    active_enterprise = {entry["model_name"].casefold() for entry in entries if entry["grade"] == "A" and entry["active"]}
    for entry in entries:
        name = entry["model_name"]
        entry["modality"] = MEDIA_MODELS.get(name, "text")
        entry["pricing_basis"] = entry["active"] and (entry["grade"] == "A" or name.casefold() not in active_enterprise)
        entry["token_price_rules"] = []
        entry["pricing_blocker"] = ""
        if not entry["pricing_basis"]:
            continue
        if entry["modality"] != "text":
            entry["pricing_blocker"] = "media quotation requires a verified media billing contract"
            continue
        share = Decimal("0.8") if entry["grade"] == "A" else Decimal("0.7")
        try:
            entry["token_price_rules"] = parse_token_price_rules(
                entry["official_price_text"], Decimal(entry["quoted_fraction"]), share
            )
        except ValueError as exc:
            entry["pricing_blocker"] = str(exc)
    manifest = {
        "workbook_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "entries": entries,
    }
    target.write_text(json.dumps(manifest, ensure_ascii=True, separators=(",", ":")) + "\n", encoding="ascii")


if __name__ == "__main__":
    main()
